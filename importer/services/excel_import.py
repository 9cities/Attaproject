from openpyxl import load_workbook


class ExcelImporter:

    def __init__(self, file):
        self.file = file
        self.workbook = None
        self.sheet = None

    def open(self):
        self.workbook = load_workbook(self.file)
        self.sheet = self.workbook.active

    def headers(self):
        return [
            str(cell.value).strip() if cell.value is not None else ""
            for cell in self.sheet[1]
        ]

    def rows(self):
        headers = self.headers()
        for row in self.sheet.iter_rows(
            min_row=2,
            values_only=True
        ):
            yield {
                headers[idx]: value
                for idx, value in enumerate(row)
                if idx < len(headers)
            }
